import os
import tempfile
from pathlib import Path

from langchain_community.document_loaders import PyPDFLoader
from langchain_core.documents import Document

SUPPORTED_TYPES = ["pdf", "docx", "txt"]


# ------------------------------------------------------------
# PDF
# ------------------------------------------------------------
def _load_pdf(path: str, name: str):
    docs = []
    for d in PyPDFLoader(path).load():
        if d.page_content.strip():
            d.metadata = {
                "source": name,
                "page": int(d.metadata.get("page", 0)) + 1,  # 1-indexed
            }
            docs.append(d)
    return docs


# ------------------------------------------------------------
# WORD (.docx) - keeps paragraphs, headings AND tables in order
# ------------------------------------------------------------
def _load_docx(path: str, name: str):
    from docx import Document as DocxDocument
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = DocxDocument(path)
    parts = []

    for child in doc.element.body.iterchildren():
        if child.tag.endswith("}p"):
            para = Paragraph(child, doc)
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name or "") if para.style is not None else ""
            if style.startswith("Heading") or style == "Title":
                text = f"## {text}"
            parts.append(text)

        elif child.tag.endswith("}tbl"):
            table = Table(child, doc)
            for row in table.rows:
                cells = []
                for cell in row.cells:
                    t = cell.text.strip().replace("\n", " ")
                    # merged cells repeat the same text
                    if t and (not cells or cells[-1] != t):
                        cells.append(t)
                if cells:
                    parts.append(" | ".join(cells))

    text = "\n\n".join(parts)
    if not text.strip():
        return []
    return [Document(page_content=text, metadata={"source": name})]


# ------------------------------------------------------------
# TXT - tries several encodings
# ------------------------------------------------------------
def _load_txt(path: str, name: str):
    raw = Path(path).read_bytes()
    text = None
    for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except (UnicodeDecodeError, UnicodeError):
            continue
    if not text or not text.strip():
        return []
    return [Document(page_content=text, metadata={"source": name})]


_LOADERS = {".pdf": _load_pdf, ".docx": _load_docx, ".txt": _load_txt}


# ------------------------------------------------------------
# Entry point used by app.py
# ------------------------------------------------------------
def load_uploaded_file(uploaded_file):
    """Returns a list of Documents (metadata['source'] = original filename)."""
    suffix = Path(uploaded_file.name).suffix.lower()

    if suffix == ".doc":
        raise ValueError(
            "Old .doc format is not supported. Please save it as .docx."
        )
    if suffix not in _LOADERS:
        raise ValueError(f"Unsupported file type: {suffix}")

    with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
        tmp.write(uploaded_file.getbuffer())
        temp_path = tmp.name

    try:
        return _LOADERS[suffix](temp_path, uploaded_file.name)
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)
